import base64
import copy
import hashlib
import io
import json
import re
from uuid import uuid4
from .domain import InputError, assess, check_chain, aggregate, check, validate_scenario
from .parser import parse, mutate, PARSER_VERSION
from .store import now
from .model import Model, ModelError, PROMPT_VERSION


FIXTURE_SOURCE='DOC-CRITERIA'


def normalize(s):
    return re.sub(r'\s+',' ',re.sub(r'-\s*\n\s*','',s)).strip()


class Service:
    def __init__(self,store,sample):
        self.store=store;self.sample=sample
        from .automation import Automation
        from .certification import Certification
        self.automation=Automation(store);self.certification=Certification(store)
        from pathlib import Path
        from .library import Library
        from .review import Review
        # Review migrates after the library so legacy requirements point at migrated version 1 documents.
        self.review=Review(store)
        self.library=Library(store,Path(store.path).parent/'documents',on_new_version=self.review.apply_document_change)
        with store.lock:self.library.migrate();self.review.migrate()
        from .portal import Portal
        self.portal=Portal(store,self.certification,self.automation)

    # Interface preferences only (adapted from codex/auditflow-prototype). Records and approvals are untouched.
    def settings(self):
        try:stored=self.store.get('meta','ui_settings')
        except KeyError:stored={}
        return {'language':stored.get('language','de'),'theme':stored.get('theme','light')}

    def save_settings(self,body):
        old=self.settings();language=body.get('language',old['language']);theme=body.get('theme',old['theme'])
        if language not in ['de','en'] or theme not in ['light','dark','system']:raise InputError('Ungültige Sprache oder Darstellung.')
        self.store.put('meta',{'id':'ui_settings','language':language,'theme':theme})
        self.store.event('ui_settings',{'language':language,'theme':theme})
        return self.settings()

    def verify(self,req):
        doc=self.store.get('document',req['source_id'])
        sec=next((x for x in doc['sections'] if x['id']==req['section_id']),None)
        req['citation_status']='verified' if sec and normalize(req['source_excerpt']) in normalize(sec['text']) and normalize(req['source_excerpt']) else 'mismatch'
        return req

    def approve(self,kind,id,body):
        if kind not in ['requirement','case']:raise InputError('Unbekannte Freigabeart.')
        if kind=='requirement':return self.review.approve({**body,'id':id})
        if not body.get('reviewer','').strip() or not body.get('note','').strip():
            raise InputError('Prüfername und fachliche Begründung sind erforderlich.')
        item=self.store.get(kind,id);approved=body.get('approve') is True
        if kind=='requirement':
            self.verify(item)
            if approved and item['citation_status']!='verified':raise InputError('Quellenzitat stimmt nicht überein.')
            item['status']='approved' if approved else 'rejected';item['semantic_review']=item['status']
        else:
            if approved and item['expected_message'] is None:raise InputError('Offene Erwartung kann nicht freigegeben werden.')
            item['review_status']='approved' if approved else 'rejected'
        item['reviewer']=body['reviewer'].strip();item['review_note']=body['note'].strip();item['reviewed_at']=now()
        self.store.put(kind,item);self.store.event('review',{'kind':kind,'id':id,'approved':approved,'reviewer':item['reviewer'],'note':item['review_note']})
        return item

    def import_document(self,body):
        name=body.get('name','').strip()
        if not name or len(name)>200:raise InputError('Dateiname fehlt oder ist zu lang.')
        try:raw=base64.b64decode(body['base64'],validate=True)
        except (KeyError,ValueError):raise InputError('Ungültiger Datei-Upload.') from None
        if len(raw)>5_000_000:raise InputError('Maximal 5 MB je Dokument.')
        sections=[]
        suffix=name.lower().rsplit('.',1)[-1]
        page_range=None
        if body.get('pages_from') is not None or body.get('pages_to') is not None:
            if suffix!='pdf':raise InputError('Eine Seitenauswahl ist nur bei PDF möglich.')
            start,end=body.get('pages_from'),body.get('pages_to')
            if type(start) is not int or type(end) is not int or not 1<=start<=end or end-start>=100:
                raise InputError('PDF-Seitenbereich: positive ganze Zahlen, höchstens 100 Seiten, von ≤ bis.')
            page_range={'from':start,'to':end}
        if suffix in ['txt','md']:
            try:text=raw.decode('utf-8-sig')
            except UnicodeError:raise InputError('Textdatei muss UTF-8 sein.') from None
            parts=re.split(r'\n\s*\n',text)
            sections=[{'id':f'p{i+1}','location':f'Abschnitt {i+1}','text':s} for i,s in enumerate(parts) if s.strip()]
        elif suffix=='pdf':
            try:
                import pdfplumber
                with pdfplumber.open(io.BytesIO(raw)) as pdf:
                    if len(pdf.pages)>500:raise InputError('Maximal 500 PDF-Seiten.')
                    if page_range and page_range['to']>len(pdf.pages):
                        raise InputError(f'Der Seitenbereich liegt außerhalb des PDF ({len(pdf.pages)} Seiten).')
                    selected=range(page_range['from']-1,page_range['to']) if page_range else range(len(pdf.pages))
                    sections=[{'id':f'page-{i+1}','page':i+1,'location':f'PDF-Seite {i+1}','text':pdf.pages[i].extract_text() or ''} for i in selected]
            except ImportError:raise InputError('PDF-Import benötigt die optionalen Pakete aus requirements.txt.') from None
        elif suffix=='docx':
            try:
                from zipfile import ZipFile
                with ZipFile(io.BytesIO(raw)) as z:
                    if sum(x.file_size for x in z.infolist())>30_000_000:raise InputError('Entpacktes DOCX ist zu groß.')
                from docx import Document
                doc=Document(io.BytesIO(raw))
                sections=[{'id':f'p{i+1}','location':f'Absatz {i+1}','text':p.text} for i,p in enumerate(doc.paragraphs) if p.text.strip()]
                sections += [{'id':f't{ti+1}r{ri+1}','location':f'Tabelle {ti+1}, Zeile {ri+1}','text':' | '.join(c.text for c in r.cells)} for ti,t in enumerate(doc.tables) for ri,r in enumerate(t.rows)]
            except ImportError:raise InputError('DOCX-Import benötigt python-docx.') from None
        else:raise InputError('Unterstützt: TXT, MD, PDF, DOCX.')
        if not sections or not any(x['text'].strip() for x in sections):raise InputError('Kein Text erkannt. Für Scans ist externe OCR erforderlich.')
        hash=hashlib.sha256(raw).hexdigest()
        existing=next((d for d in self.store.list('document') if d['sha256']==hash and d.get('version')==(body.get('version') or 'unbekannt') and d.get('valid_from')==(body.get('valid_from') or None) and d.get('page_range')==page_range),None)
        if existing:return existing
        doc={'id':str(uuid4()),'name':name,'version':body.get('version') or 'unbekannt','valid_from':body.get('valid_from') or None,
             'sha256':hash,'sections':sections,'kind':suffix,'page_range':page_range,
             'note':('Teilimport: Originalseiten '+str(page_range['from'])+'–'+str(page_range['to'])+'. ' if page_range else '')+'Importiert; Gültigkeit und Vollständigkeit prüfen.'}
        self.store.put('document',doc);self.store.event('document_import',{'id':doc['id'],'name':name,'hash':hash})
        return doc

    def draft(self,body):
        doc=self.store.get('document',body['document_id'])
        sec=next((s for s in doc['sections'] if s['id']==body['section_id']),None)
        if not sec:raise InputError('Quellabschnitt nicht gefunden.')
        if len(sec['text'])>35000:raise InputError('Pilotabschnitt zu lang; kleineren Abschnitt auswählen.')
        live=body.get('mode')=='live'
        if live:
            drafts=Model().extract(doc,sec)
        else:
            # Extraction assistance available with zero model calls: explicitly literal candidates.
            sentences=[s.strip() for s in re.split(r'\n|(?<=[.;])\s+',sec['text']) if len(s.strip())>35][:8]
            drafts=[{'title':s[:85],'criterion':s,'source_excerpt':s,'conditions':[],'exceptions':[],
                     'questions':['Wörtlicher Offline-Kandidat: Kontext, Ausnahmen und Regelzuordnung fachlich ergänzen.']} for s in sentences]
        created=[]
        for item in drafts:
            r={**item,'id':str(uuid4()),'source_id':doc['id'],'section_id':sec['id'],'status':'pending',
               'semantic_review':'pending','mapping':None,'ruleset':'unmapped','citation_status':'unverified',
               'model_id':Model().model if live else 'offline-literal','prompt_version':PROMPT_VERSION,
               'note':'KI-Vorschlag' if live else 'Wörtlicher Offline-Vorschlag, keine KI-Ausführung'}
            self.verify(r);self.store.put('requirement',r);created.append(self.review.register(r['id'],'ki_entwurf' if live else 'offline_literal'))
        self.store.event('requirements_drafted',{'document_id':doc['id'],'mode':'live' if live else 'offline','ids':[r['id'] for r in created]})
        return created

    def _approved_source(self,rule):
        try:return self.store.get('requirement_version',f"{rule['id']}@v{rule['approved_version_no']}")['source_id']
        except (KeyError,TypeError):return rule.get('source_id')

    def _versions(self,ids):
        reqs=self.store.get_many('requirement',list(ids))
        return {id:reqs[id].get('version_no',1) for id in ids if id in reqs}

    def requirement_versions(self,id):
        req=self.store.get('requirement',id)
        versions=sorted(self.store.filtered('requirement_version','requirement_id',id),key=lambda v:v['version_no'])
        return {'current':req,'approved_version_no':req.get('approved_version_no'),'versions':versions}

    def create_case(self,body):
        s=body['scenario'];validate_scenario(s)
        name=body.get('name','').strip()
        if not name:raise InputError('Testfallname fehlt.')
        ruleset=body.get('ruleset','fixture-2026')
        if ruleset!='fixture-2026':raise InputError('Nur Fixture-Regelsatz ist ausführbar; andere Versionen sind nicht implementiert.')
        expected=body.get('expected_message')
        if expected is not None and not isinstance(expected,bool):raise InputError('Erwartung muss true, false oder null sein.')
        test_type=body.get('test_type','custom')
        if test_type not in ['custom','positive','negative','boundary','missing_data']:raise InputError('Unbekannte Testfallart.')
        item={'id':str(uuid4()),'name':name,'scenario':s,'expected_message':expected,'ruleset':ruleset,'test_type':test_type,
              'review_status':'pending','expectation_source':body.get('expectation_source','Manueller Entwurf; Quelle ergänzen.'),
              'requirement_ids':['C1','C2','C4'],'fixture_questions':[],'expectation_edited':False,'outdated':None,'regenerated_from':None}
        item['requirement_versions']=self._versions(item['requirement_ids'])
        self.store.put('case',item);self.store.event('case_created',{'id':item['id']});return item

    def generate_cases(self,body):
        from datetime import timedelta
        from .domain import iso, add_months
        baseline=self.store.get('case',body.get('case_id','CASE-001'))
        if baseline['ruleset']!='fixture-2026':raise InputError('Kein ausführbarer Vorlagenstand für diesen Ausgangsfall.')
        rules=[self.store.get('requirement',id) for id in ['C1','C2','C4']]
        if not all(self.review.effective(r) for r in rules):
            raise InputError('Zuerst C1, C2 und C4 fachlich freigeben.')
        mode=body.get('mode','boundary');drafts=[]
        test_type=body.get('test_type','boundary')
        if test_type not in ['positive','negative','boundary','missing_data']:raise InputError('Unbekannte Testfallart.')
        if mode=='live':
            reply=Model().call(
                'Create at most 4 synthetic DBVO demo test-case DRAFTS from the approved demo rules. '
                'Do not claim legal validity or approve expectations. Return ONLY JSON: '
                '{"cases":[{"name":"...","scenario":{"current":{"begin":"YYYY-MM-DD","end":"YYYY-MM-DD",'
                '"certified":true},"prior":[{"begin":"YYYY-MM-DD","end":"YYYY-MM-DD","certified":true}],'
                '"report_date":"YYYY-MM-DD"},"expected_message":true,"expectation_source":"draft rationale"}]}. '
                'Expected message may be true, false or null. Explain uncertain outcomes. No real personal data.',
                json.dumps({'rules':rules,'baseline':baseline['scenario'],'test_type':test_type,
                            'instruction':'Generate this requested test type only. Unresolved boundaries stay null; missing data is not a confirmed rejection.'},ensure_ascii=False))
            try:
                reply=re.sub(r'^```(?:json)?\s*|\s*```$','',reply.strip())
                drafts=json.loads(reply)['cases']
                if not isinstance(drafts,list) or not 1<=len(drafts)<=4:raise ValueError()
                for d in drafts:
                    if not isinstance(d,dict) or not isinstance(d.get('name'),str) or not d['name'].strip():raise ValueError()
                    validate_scenario(d['scenario'])
                    if d.get('expected_message') is not None and not isinstance(d['expected_message'],bool):raise ValueError()
                    d['expectation_source']='KI-Entwurf, unabhängig zu prüfen: '+str(d.get('expectation_source',''))
            except (KeyError,ValueError,TypeError):
                raise ModelError('Ungültiger Testfallentwurf; keine KI-Testfälle gespeichert.') from None
        elif mode=='template':
            # Authored synthetic fixtures; no legal rules are inferred from arbitrary text.
            scenario={'current':{'begin':'2026-06-01','end':'2026-06-10','certified':True},
                      'prior':[{'begin':'2026-05-01','end':'2026-05-25','certified':True}],
                      'report_date':'2026-06-10'}
            if test_type=='boundary':return self.generate_cases({**body,'mode':'boundary'})
            if test_type=='negative':scenario['current']['certified']=False
            if test_type=='missing_data':scenario['current']['certified']=None
            expected={'positive':True,'negative':False,'missing_data':None}[test_type]
            drafts=[{'name':{'positive':'Positivfall: attestierte Anfrage','negative':'Negativfall: Anfrage unterdrücken','missing_data':'Fehlende Daten: Nachweisstatus unbekannt'}[test_type],
                     'scenario':scenario,'expected_message':expected,
                     'expectation_source':'Synthetische Vorlage, 35 Tage im Demobestand. Erwartung und fachlichen Kontext unabhängig prüfen.'}]
        elif mode=='boundary':
            start=iso(baseline['scenario']['current']['begin'])
            for n in [29,30,31]:
                scenario=copy.deepcopy(baseline['scenario'])
                scenario['current']['end']=str(start+timedelta(days=9))
                a=add_months(start.replace(day=1),-1)
                scenario['prior']=[{'begin':str(a),'end':str(a+timedelta(days=n-11)),'certified':True}]
                scenario['current']['certified']=True
                drafts.append({'name':f'Generierter Grenzfall: {n} Tage','scenario':scenario,'expected_message':n>=30,
                               'expectation_source':'Deterministischer synthetischer Entwurf; 10 aktuelle Tage plus '+str(n-10)+' frühere Tage. Fachlich prüfen.'})
            scenario=copy.deepcopy(baseline['scenario']);scenario['current']['end']=None;scenario.pop('open_end_policy',None)
            drafts.append({'name':'Generierter Grenzfall: offene AU','scenario':scenario,'expected_message':None,
                           'expectation_source':'Zählkonvention nicht bestätigt; Erwartung bleibt offen.'})
        else:raise InputError('Unbekannter Generierungsmodus.')
        created=[self.create_case({**d,'ruleset':baseline['ruleset'],'test_type':test_type if mode!='boundary' else 'boundary'}) for d in drafts]
        for item in created:
            item.update(generation_mode=mode,model_id=Model().model if mode=='live' else None,
                        prompt_version='test-drafts-v2',requirement_ids=[r['id'] for r in rules],
                        requirement_versions=self._versions([r['id'] for r in rules]),requirements_snapshot=copy.deepcopy(rules))
            self.store.put('case',item)
        self.store.event('cases_generated',{'mode':mode,'baseline':baseline['id'],'ids':[c['id'] for c in created]})
        return created

    def run(self,body):
        case=self.store.get('case',body.get('case_id','CASE-001'))
        if case['ruleset']!='fixture-2026':raise InputError('Kein ausführbarer Prüfer für diesen Regelstand.')
        mode=body.get('mode','scenario')
        if mode not in ['scenario','message']:raise InputError('Unbekannter Prüfmodus.')
        rules={r['id']:r for r in self.store.list('requirement')}
        results=assess(case['scenario']);comparison=[];parsed=None;error=None
        raw=None
        if mode=='message':
            try:
                if body.get('base64'):
                    raw=base64.b64decode(body['base64'],validate=True)
                elif body.get('mutation'):
                    raw=mutate(self.sample,body['mutation'])
                else:raw=self.sample
                parsed=parse(raw)
                if parsed['abgabegrund']!='41' or 'dbvo' not in parsed:raise InputError('Ausführbarer Fixture-Prüfer unterstützt nur Abgabegrund 41 mit DBVO.')
                c=check_chain(parsed['dbvo']['periods'],case['scenario']);results.append(c)
                results.append(check('FORMAT','pass','Konsistentes Fixture-Layout',parsed['byte_length'],'raw_message'))
                results.append(check('GROUND_STATE','pass' if parsed['dbvo']['insurer_fields_zero'] else 'fail',
                                     'KK-Felder in AG-Fixture in Grundstellung',parsed['dbvo']['insurer_fields_zero'],'message.DBVO'))
                # All comparisons are explicitly fixture expectations; they do not certify the standard.
                expected={'reason':'41','current_begin':case['scenario']['current']['begin'],
                          'periods':sorted([{'begin':p['begin'],'end':p['end']} for p in case['scenario']['prior'] if p.get('certified')],key=lambda p:p['begin'],reverse=True)}
                actual={'reason':parsed['abgabegrund'],'current_begin':parsed['dbvo']['au_ab_ag'],'periods':parsed['dbvo']['periods']}
                if case['id']=='CASE-001':
                    expected.update(datensatz_id='00000000-0000-4000-8000-000000000001',azvu='DEMO-REFERENCE-001')
                    actual.update(datensatz_id=parsed['datensatz_id'],azvu=parsed.get('dbid',{}).get('aktenzeichen_verursacher'))
                for field in expected:
                    comparison.append({'field':field,'expected':expected[field],'actual':actual.get(field),
                                       'status':'pass' if expected[field]==actual.get(field) else 'fail'})
                if case['expected_message'] is False:
                    comparison.append({'field':'message_presence','expected':False,'actual':True,'status':'fail'})
            except (InputError,ValueError) as e:error=str(e)
        else:
            # Internal simulator shares the rule engine: explicitly not independent vendor validation.
            statuses=[r['status'] for r in results]
            simulated=False if 'fail' in statuses else None if 'not_assessable' in statuses else True
            if body.get('software_mode')=='buggy':simulated=True
            expected=case['expected_message']
            comparison=[{'field':'message_presence','expected':expected,'actual':simulated,
                         'status':'not_assessable' if expected is None or simulated is None else 'pass' if expected==simulated else 'fail'}]
        # Substantive approval is distinct from arithmetic and citation verification.
        for result in results:
            rule=rules.get(result['requirement_id'])
            if rule:
                result['source']={'document_id':rule['source_id'],'section_id':rule['section_id'],'excerpt':rule['source_excerpt']}
                result['computed_status']=result['status']
                if not self.review.effective(rule):
                    result['status']='not_assessable';result['note']+=' Fachregel noch nicht freigegeben.'
                elif self._approved_source(rule)!=FIXTURE_SOURCE:
                    # The fixture checker implements the synthetic v1 rules only; never apply it to a newer source version.
                    result['status']='not_assessable'
                    result['note']+=f" Freigegebene Version {rule.get('version_no')} beruht auf einem neueren Quellstand; der Fixture-Prüfer implementiert nur {FIXTURE_SOURCE}."
            elif result['requirement_id'] in ['C1','C2','C3','C4']:
                # The requirement was deleted together with its source document: never count it as passed.
                result['computed_status']=result['status'];result['status']='not_assessable';result['note']+=' Anforderung wurde mit ihrem Quelldokument gelöscht.'
            elif result['requirement_id'] in ['FORMAT','GROUND_STATE']:
                result['source']={'document_id':'fixture','section_id':'inferred-layout','excerpt':'Sample-specific structural check; normative technical schema not verified.'}
            if parsed and result['requirement_id']=='C3':
                block=parsed['blocks']['DBVO'];result['evidence'].update(block='DBVO',position=block['position'],position_basis='byte-1-based',raw=block['raw'])
        if not self.review.case_usable(case):
            for row in comparison:
                row['computed_status']=row['status'];row['status']='not_assessable'
        # Compliance of a message vs eligibility of a scenario are different result axes.
        domain_status=aggregate(results)
        comparison_status=aggregate(comparison)
        combined=results+comparison
        run={'id':str(uuid4()),'created_at':now(),'case_id':case['id'],'case_name':case['name'],'ruleset':case['ruleset'],
             'mode':mode,'execution_mode':'simulation' if mode=='scenario' else 'imported_fixture',
             'overall_status':aggregate(combined if mode=='message' else comparison + [dict(status='not_assessable') for r in results if r['status']=='not_assessable'], bool(error)),'rule_status':domain_status,'comparison_status':comparison_status,
             'results':results,'comparison':comparison,'parsed':parsed,'error':error,
             'snapshot':copy.deepcopy(case),'requirements_snapshot':[rules[k] for k in ['C1','C2','C3','C4'] if k in rules],
             'checker_version':'fixture-rules-0.1','parser_version':PARSER_VERSION,'input_hash':hashlib.sha256(raw or json.dumps(case['scenario'],sort_keys=True).encode()).hexdigest(),
             'limitations':['Keine echte Abrechnungssoftware angebunden.','Nur vier Beispielkriterien; keine vollständige EEL-Konformitätsprüfung.',
                            'Fixture-Layout abgeleitet; technische Anlage nicht verifiziert.','Kein Zertifikat und keine verbindliche Leistungsberechnung.',
                            'Simulationsentscheidung und Regelprüfung teilen Logik; unabhängige Herstellerprüfung nicht nachgewiesen.']}
        self.store.put('run',run);self.store.event('test_run',{'run_id':run['id'],'status':run['overall_status'],'input_hash':run['input_hash']})
        return run

    def coverage(self):
        reqs=[r for r in self.store.list('requirement') if self.review.effective(r)]
        cases=self.store.list('case');runs=self.store.list('run');cells=[]
        for r in reqs:
            for c in cases:
                if r['id'] not in c['requirement_ids']:continue
                latest=next((v for v in reversed(runs) if v['case_id']==c['id'] and v['ruleset']==c['ruleset'] and any(x['requirement_id']==r['id'] for x in v['results'])),None)
                result=next((x for x in latest['results'] if x['requirement_id']==r['id']),None) if latest else None
                stale=bool(latest and (latest['snapshot']!=c or next((x for x in latest['requirements_snapshot'] if x['id']==r['id']),None)!=r))
                cells.append({'requirement_id':r['id'],'case_id':c['id'],'case_approved':self.review.case_usable(c),
                              'status':'stale' if stale else result['status'] if result else 'not_run'})
        gaps=[r['id'] for r in reqs if not any(c['requirement_id']==r['id'] and c['case_approved'] for c in cells)]
        return {'requirements':reqs,'cases':cases,'cells':cells,'gaps':gaps}

    def diff(self,body):
        old=self.store.get('document',body['old_id']);new=self.store.get('document',body['new_id'])
        if old['id']==new['id']:raise InputError('Zwei unterschiedliche Dokumentstände auswählen.')
        a={s['id']:s for s in old['sections']};b={s['id']:s for s in new['sections']};changes=[]
        for id in sorted(a.keys()|b.keys()):
            if id not in a:kind='added'
            elif id not in b:kind='removed'
            elif normalize(a[id]['text'])!=normalize(b[id]['text']):kind='changed'
            else:continue
            changes.append({'section_id':id,'kind':kind,'before':a.get(id,{}).get('text'),'after':b.get(id,{}).get('text')})
        ids={x['section_id'] for x in changes}
        reqs=[r for r in self.store.list('requirement') if r['source_id']==old['id'] and any(r['section_id']==id or r['section_id'].startswith(id+'.') for id in ids)]
        affected={r['id'] for r in reqs}
        return {'changes':changes,'requirements':reqs,'cases':[c for c in self.store.list('case') if affected.intersection(c['requirement_ids'])],
                'limitations':'Struktureller Abschnittsvergleich; gleiche Abschnitts-IDs vorausgesetzt. Verschobene Definitionen und indirekte Abhängigkeiten manuell prüfen.'}
