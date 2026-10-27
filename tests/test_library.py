import base64
import io
import json
import os
import tempfile
import unittest
from importlib.util import find_spec
from pathlib import Path
from auditflow.store import Store
from auditflow.library import Library, diff_paragraphs, build_paragraphs, DOC_TYPES
from auditflow.domain import InputError
ROOT=Path(__file__).resolve().parents[1]
B64=lambda s:base64.b64encode(s.encode() if isinstance(s,str) else s).decode()
# Fiktive Testtexte, keine echte Rechtsquelle.
V1='Fiktive Prüfkriterien für Testzwecke.\n\nRegel A: Die Anfrage ist nur bei bestätigter AU zulässig.\n\nRegel B: Der Abstand wird in Kalendermonaten berechnet.\n\nRegel C: Die Antwort enthält das Aktenzeichen.'
V2='Fiktive Prüfkriterien für Testzwecke.\n\nRegel A: Die Anfrage ist nur bei bestätigter und aktueller AU zulässig.\n\nRegel C: Die Antwort enthält das Aktenzeichen.\n\nRegel D: Stornierungen werden protokolliert.'

class LibraryTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();base=Path(self.temp.name)
        seed=json.loads((ROOT/'data/seed.json').read_text(encoding='utf-8'))
        self.store=Store(base/'db.sqlite3',seed);self.files=base/'files';self.calls=[]
        def hook(diff):
            self.calls.append(diff['id']);diff=dict(diff);diff['affected_requirement_ids']=['C1']
            diff['changes']=[dict(c,classification='fachlich',classification_source='ki') for c in diff['changes']];return diff
        self.lib=Library(self.store,self.files,on_new_version=hook);self.lib.migrate()
    def tearDown(self):self.temp.cleanup()
    def upload(self,text,name='Pruefkriterien 2026.txt',label='2026-01',library_id=None,**extra):
        a=self.lib.analyse({'name':name,'base64':B64(text),'version_label':label,**extra})
        self.assertIsNone(a['duplicate'])
        return a,self.lib.confirm({'upload_id':a['upload_id'],'library_id':library_id,'doc_type':'kriterien'})

    def test_new_document_is_version_one_and_file_stored(self):
        a,r=self.upload(V1)
        self.assertIsNone(a['suggestion']);self.assertEqual(a['paragraphs'],4)
        doc=r['document'];self.assertEqual((doc['version_no'],doc['version'],doc['kind'],doc['page_range']),(1,'2026-01','txt',None))
        self.assertTrue(doc['file_stored']);self.assertIsNone(r['diff'])
        self.assertEqual(doc['sections'][1]['id'],'p2');self.assertEqual(doc['paragraphs'][1]['ref'],'¶2')
        self.assertEqual([s['label'] for s in r['steps']],['Datei geprüft','Text gelesen','Dokument erkannt','Gespeichert','Änderungen verglichen','Anforderungen aktualisiert'])
        raw,mime,name=self.lib.file(doc['id']);self.assertEqual(raw,V1.encode());self.assertTrue(mime.startswith('text/plain'))
        self.assertEqual(self.lib.text(doc['id'])['paragraphs'][0]['location'],'Absatz 1')
        actions=[e['action'] for e in self.store.events()];self.assertIn('library_upload',actions);self.assertIn('library_confirm',actions)

    def test_duplicate_rejected_even_with_other_label(self):
        _,r=self.upload(V1);before=len(self.store.list('upload'))
        a=self.lib.analyse({'name':'anders.txt','base64':B64(V1),'version_label':'2099-x','valid_from':'2027-01-01'})
        self.assertEqual(a['duplicate']['document_id'],r['document']['id']);self.assertIsNone(a['upload_id'])
        self.assertIn('bereits als',a['duplicate']['message']);self.assertEqual(len(self.store.list('upload')),before)

    def test_similar_file_suggested_and_new_version_keeps_old(self):
        _,r1=self.upload(V1);lib_id=r1['library_doc']['id']
        a=self.lib.analyse({'name':'Pruefkriterien 2027.txt','base64':B64(V2),'version_label':'2027-01'})
        self.assertEqual(a['suggestion']['library_id'],lib_id);self.assertGreaterEqual(a['suggestion']['score'],0.45)
        r2=self.lib.confirm({'upload_id':a['upload_id'],'library_id':lib_id})
        self.assertEqual(r2['document']['version_no'],2);self.assertEqual(len(r2['library_doc']['version_ids']),2)
        self.assertEqual(self.store.get('document',r1['document']['id'])['version'],'2026-01')
        with self.assertRaises(InputError):self.lib.confirm({'upload_id':a['upload_id'],'library_id':lib_id})
        st=self.lib.version_status(self.store.get('document',r1['document']['id']),'2026-10-27')
        self.assertEqual(st['label'],'Ersetzt durch Version 2')
        folder=next(f for f in self.lib.listing()['folders'] if f['doc_type']=='kriterien')
        entry=next(d for d in folder['documents'] if d['library_id']==lib_id)
        self.assertEqual(entry['current']['version_no'],2);self.assertEqual([v['version_no'] for v in entry['versions']],[2,1])

    def test_diff_kinds_word_ops_and_hook_persisted(self):
        _,r1=self.upload(V1);_,r2=self.upload(V2,name='Kriterien neu.txt',label='2027-01',library_id=r1['library_doc']['id'])
        d=r2['diff'];kinds={c['kind'] for c in d['changes']}
        self.assertEqual(kinds,{'added','removed','changed'})
        self.assertEqual((d['summary']['added'],d['summary']['removed'],d['summary']['changed']),(1,1,1))
        ch=next(c for c in d['changes'] if c['kind']=='changed')
        self.assertIn({'op':'insert','text':'und aktueller '},ch['words'])
        self.assertEqual((ch['old_ref'],ch['new_ref']),('¶2','¶2'))
        self.assertEqual(self.calls,[d['id']])
        stored=self.lib.diff(d['id']);self.assertEqual(stored['affected_requirement_ids'],['C1'])
        self.assertEqual(stored['summary']['fachlich'],3);self.assertEqual(stored['summary']['affected'],1)
        self.assertEqual(self.store.get('document',r2['document']['id'])['diff_id'],d['id'])
        self.assertEqual(r2['steps'][-1]['status'],'ok');self.assertIn('1 Anforderungen betroffen',r2['summary'])
        self.assertEqual(self.lib.version(r2['document']['id'])['diff']['id'],d['id'])

    def test_initial_classification_without_hook(self):
        lib=Library(self.store,self.files)
        a=lib.analyse({'name':'x.md','base64':B64(V1),'version_label':'1'});r1=lib.confirm({'upload_id':a['upload_id']})
        a=lib.analyse({'name':'x.md','base64':B64(V2),'version_label':'2'});r2=lib.confirm({'upload_id':a['upload_id'],'library_id':r1['library_doc']['id']})
        self.assertTrue(all(c['classification']=='offen' and c['classification_source']=='keine' for c in r2['diff']['changes']))
        self.assertEqual(r2['steps'][-1]['status'],'skipped')

    def test_page_shift_is_not_a_change(self):
        old=build_paragraphs([{'id':'page-1','page':1,'location':'PDF-Seite 1','text':'Einleitung.\n\nRegel A gilt immer.'},
                              {'id':'page-2','page':2,'location':'PDF-Seite 2','text':'Regel B gilt nie.'}])
        new=build_paragraphs([{'id':'page-1','page':1,'location':'PDF-Seite 1','text':'Einleitung.\n\nNeuer Vorspann ist eingefügt.'},
                              {'id':'page-2','page':2,'location':'PDF-Seite 2','text':'Regel A gilt immer.'},
                              {'id':'page-3','page':3,'location':'PDF-Seite 3','text':'Regel B gilt nie.'}])
        changes=diff_paragraphs(old,new)
        self.assertEqual([c['kind'] for c in changes],['added']);self.assertEqual(changes[0]['new_location'],'Seite 1, Absatz 2')
        # Ein über die Seitengrenze getrennter Satz zählt als ein Absatz.
        split=build_paragraphs([{'id':'page-1','page':1,'location':'PDF-Seite 1','text':'Regel A gilt'},{'id':'page-2','page':2,'location':'PDF-Seite 2','text':'immer.'}])
        self.assertEqual(diff_paragraphs(split,[{'ref':'¶1','page':1,'location':'x','text':'Regel A gilt immer.'}]),[])
        # Reihenfolge getauscht, Text gleich: keine Änderung.
        self.assertEqual(diff_paragraphs(old,list(reversed(old))),[])

    def test_date_suggestions(self):
        a=self.lib.analyse({'name':'d.txt','base64':B64('Diese fiktive Fassung gilt ab 01.01.2027 und ist gültig bis 2027-12-31.\n\nSie tritt am 1.7.2027 in Kraft.'),'version_label':'x'})
        got={(d['field'],d['value']) for d in a['date_suggestions']}
        self.assertEqual(got,{('valid_from','2027-01-01'),('valid_until','2027-12-31'),('valid_from','2027-07-01')})
        self.assertEqual(self.lib.suggest_dates('gilt ab 31.02.2027'),[])
        r=self.lib.confirm({'upload_id':a['upload_id']})
        self.assertIsNone(r['document']['valid_from']);self.assertFalse(r['document']['validity_confirmed'])

    def test_version_status_rules_and_withdraw(self):
        _,r1=self.upload(V1,valid_until='2026-06-30');lib_id=r1['library_doc']['id'];d1=r1['document']['id']
        self.assertEqual(self.lib.version_status(self.store.get('document',d1),'2026-05-01')['label'],'Bereit')
        self.assertEqual(self.lib.version_status(self.store.get('document',d1),'2026-07-01')['label'],'Abgelaufen')
        a=self.lib.analyse({'name':'n.txt','base64':B64(V2),'version_label':'2027','valid_from':'2027-01-01'})
        r2=self.lib.confirm({'upload_id':a['upload_id'],'library_id':lib_id});d2=r2['document']['id']
        self.assertTrue(r2['document']['validity_confirmed'])
        self.assertEqual(self.lib.version_status(self.store.get('document',d1),'2026-05-01')['label'],'Bereit')
        self.assertEqual(self.lib.version_status(self.store.get('document',d1),'2027-01-01')['label'],'Ersetzt durch Version 2')
        with self.assertRaises(InputError):self.lib.withdraw({'document_id':d2,'reason':'','by':'Tester'})
        w=self.lib.withdraw({'document_id':d2,'reason':'Falsche Datei','by':'Tester'})
        self.assertEqual(w['withdrawn']['reason'],'Falsche Datei')
        self.assertEqual(self.lib.version_status(w,'2027-02-01')['withdrawn'],'Zurückgezogen: Falsche Datei')
        self.assertEqual(self.lib.version_status(self.store.get('document',d1),'2026-05-01')['label'],'Bereit')
        with self.assertRaises(InputError):self.lib.withdraw({'document_id':d2,'reason':'x','by':'y'})
        # Vergleich der nächsten Version läuft gegen die letzte nicht zurückgezogene Version.
        a=self.lib.analyse({'name':'n3.txt','base64':B64(V2+'\n\nRegel E: fiktiv.'),'version_label':'2027b'})
        r3=self.lib.confirm({'upload_id':a['upload_id'],'library_id':lib_id})
        self.assertEqual((r3['document']['version_no'],r3['diff']['old_version_id']),(3,d1))
        self.assertIn('library_withdraw',[e['action'] for e in self.store.events()])

    def test_migrate_seed_idempotent_and_file_not_stored(self):
        docs={d['id']:d for d in self.store.list('document')}
        self.assertEqual(docs['DOC-CRITERIA']['version_no'],1);self.assertFalse(docs['DOC-CRITERIA']['file_stored'])
        self.assertEqual(docs['DOC-CRITERIA']['doc_type'],'kriterien');self.assertEqual(docs['DOC-SCENARIO']['doc_type'],'sonstige')
        self.assertEqual(docs['DOC-CRITERIA']['paragraphs'][0]['location'],'Selbst verfasste Demoregel C1')
        libs=len(self.store.list('library_doc'));self.assertEqual(self.lib.migrate(),[]);self.assertEqual(len(self.store.list('library_doc')),libs)
        with self.assertRaises(InputError):self.lib.file('DOC-CRITERIA')
        self.assertTrue(self.lib.text('DOC-CRITERIA')['paragraphs'])
        self.assertEqual(set(self.lib.version('DOC-CRITERIA')['requirement_ids']),{'C1','C2','C3','C4'})
        folder=next(f for f in self.lib.listing()['folders'] if f['doc_type']=='kriterien')
        self.assertEqual(folder['documents'][0]['pending_requirements'],4)
        self.assertEqual([f['label'] for f in self.lib.listing()['folders']],list(DOC_TYPES.values()))

    def test_input_validation(self):
        with self.assertRaisesRegex(InputError,'Versionsbezeichnung fehlt'):self.lib.analyse({'name':'a.txt','base64':B64('x'),'version_label':' '})
        with self.assertRaises(InputError):self.lib.analyse({'name':'a.exe','base64':B64('x'),'version_label':'1'})
        with self.assertRaises(InputError):self.lib.analyse({'name':'a.txt','base64':'###','version_label':'1'})
        with self.assertRaises(InputError):self.lib.analyse({'name':'a.txt','base64':B64('x'),'version_label':'1','valid_from':'01.01.2027'})
        with self.assertRaises(InputError):self.lib.analyse({'name':'a.txt','base64':B64(b'\xff\xfe\x00'),'version_label':'1'})
        with self.assertRaises(InputError):self.lib.confirm({'upload_id':'UP-unknown'})

    def test_path_safety(self):
        _,r=self.upload(V1,name='../../x.txt');doc=r['document']
        self.assertEqual(sorted(os.listdir(self.files)),[doc['sha256']+'.txt'])
        self.assertFalse((Path(self.temp.name).parent/'x.txt').exists())
        raw,_,name=self.lib.file(doc['id']);self.assertEqual(raw,V1.encode());self.assertNotIn('/',name)
        doc['sha256']='../../etc';self.store.put('document',doc)
        with self.assertRaises(InputError):self.lib.file(doc['id'])

    @unittest.skipUnless(find_spec('pdfplumber') and find_spec('reportlab'),'pdfplumber/reportlab nicht installiert')
    def test_pdf_pages(self):
        from reportlab.pdfgen import canvas
        buf=io.BytesIO();c=canvas.Canvas(buf)
        for t in ['Fiktive Seite eins.','Fiktive Seite zwei.']:c.drawString(72,720,t);c.showPage()
        c.save();a=self.lib.analyse({'name':'p.pdf','base64':B64(buf.getvalue()),'version_label':'1'})
        r=self.lib.confirm({'upload_id':a['upload_id']});self.assertEqual([p['page'] for p in r['document']['paragraphs']],[1,2])
        self.assertEqual(self.lib.file(r['document']['id'])[1],'application/pdf')

    @unittest.skipUnless(find_spec('docx'),'python-docx nicht installiert')
    def test_docx(self):
        from docx import Document
        d=Document();d.add_paragraph('Fiktiver Absatz eins.');d.add_paragraph('Fiktiver Absatz zwei.');buf=io.BytesIO();d.save(buf)
        a=self.lib.analyse({'name':'d.docx','base64':B64(buf.getvalue()),'version_label':'1'});self.assertEqual(a['paragraphs'],2)

if __name__=='__main__':unittest.main()
